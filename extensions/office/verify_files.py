"""Exercise actual third-party file APIs in a fresh output directory (not visual QA)."""
import argparse
import hashlib
import json
from pathlib import Path


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(output):
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=False)
    originals = {}

    doc = Document()
    doc.add_heading('项目周报', 0)
    paragraph = doc.add_paragraph()
    paragraph.add_run('已完成：').bold = True
    paragraph.add_run('读取')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '任务'; table.cell(0, 1).text = '状态'
    table.cell(1, 0).text = '写入'; table.cell(1, 1).text = '待做'
    source = output/'source.docx'; doc.save(source); originals[source.name] = digest(source)
    doc = Document(source)
    assert doc.paragraphs[1].runs[0].bold is True
    doc.paragraphs[1].runs[1].text = '读取与写入'
    doc.tables[0].cell(1, 1).text = '完成'
    doc.save(output/'edited.docx')
    doc = Document(output/'edited.docx')
    assert doc.paragraphs[1].text == '已完成：读取与写入'
    assert doc.paragraphs[1].runs[0].bold is True
    assert doc.tables[0].cell(1, 1).text == '完成'

    book = Workbook(); sheet = book.active; sheet.title = '进度'
    sheet.append(['项目', '完成量']); sheet.append(['读取', 2]); sheet.append(['写入', 3])
    sheet['B4'] = '=SUM(B2:B3)'; sheet['A1'].font = Font(bold=True)
    sheet.merge_cells('D1:E1'); sheet['D1'] = '保留合并单元格'
    source = output/'source.xlsx'; book.save(source); book.close(); originals[source.name] = digest(source)
    book = load_workbook(source, data_only=False); book['进度']['B3'] = 5
    book.save(output/'edited.xlsx'); book.close()
    book = load_workbook(output/'edited.xlsx', data_only=False); sheet = book['进度']
    assert sheet['B3'].value == 5 and sheet['B4'].value == '=SUM(B2:B3)'
    assert sheet['A1'].font.bold and 'D1:E1' in sheet.merged_cells
    book.close()
    values = load_workbook(output/'edited.xlsx', data_only=True)
    assert values['进度']['B4'].value is None, 'fixture must not pretend to calculate formulas'
    values.close()

    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = '阶段进度'
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run(); run.text = '读取完成'; run.font.bold = True
    slides.slides.add_slide(slides.slide_layouts[5]).shapes.title.text = '保留页面'
    source = output/'source.pptx'; slides.save(source); originals[source.name] = digest(source)
    slides = Presentation(source)
    slides.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].text = '读取与写入完成'
    slides.save(output/'edited.pptx'); slides = Presentation(output/'edited.pptx')
    assert len(slides.slides) == 2 and slides.slides[1].shapes.title.text == '保留页面'
    run = slides.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert run.text == '读取与写入完成' and run.font.bold

    source = output/'source.pdf'; pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    pdf = canvas.Canvas(str(source)); pdf.drawString(72, 720, 'Sumika office baseline')
    pdf.setFont('STSong-Light', 12); pdf.drawString(72, 690, '中文项目进度')
    pdf.showPage(); pdf.drawString(72, 720, 'Second page retained'); pdf.save()
    originals[source.name] = digest(source)
    reader = PdfReader(source)
    assert 'Sumika office baseline' in reader.pages[0].extract_text()
    assert '中文项目进度' in reader.pages[0].extract_text()
    writer = PdfWriter(); writer.append(reader); writer.pages[0].rotate(90)
    writer.add_metadata({'/Title': 'Updated project report'})
    writer.write(output/'edited.pdf'); writer.close()
    edited = PdfReader(output/'edited.pdf')
    assert len(edited.pages) == 2 and edited.pages[0].rotation == 90
    assert 'Second page retained' in edited.pages[1].extract_text()
    assert edited.metadata.title == 'Updated project report'

    for name, expected in originals.items():
        assert digest(output/name) == expected, f'Original changed: {name}'
    report = {'passed': True, 'checks': ['docx_create_read_edit_runs_table',
              'xlsx_create_read_edit_formula_style_merge', 'xlsx_no_fake_calculation',
              'pptx_create_read_edit_preserve_second_slide',
              'pdf_create_read_chinese_rotate_metadata', 'originals_unchanged'],
              'visual_validation': 'not_run', 'formula_calculation': 'not_run',
              'files': {p.name: digest(p) for p in sorted(output.iterdir())}}
    (output/'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(verify(args.output), ensure_ascii=False))
