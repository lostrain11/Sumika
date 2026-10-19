"""Synthetic structured Office roundtrips and real render/recalculation evidence.

Run through the enabled sumika-office Skill. PNGs require human visual review;
this deterministic fixture does not assess autonomous natural-language quality.
"""
import argparse
from datetime import date
import hashlib
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extensions.office.render import convert
from extensions.desktop.ocr.ocr import recognize
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfReader
import pypdfium2 as pdfium


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def exercise(root, soffice):
    checks = []
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(11)
    doc.sections[0].header.paragraphs[0].text = 'Sumika Office 验收页眉'
    doc.sections[0].footer.paragraphs[0].text = '合成测试材料 / 不含真实用户数据'
    doc.add_heading('项目周报：复杂结构', 0)
    p = doc.add_paragraph()
    p.add_run('保留加粗：').bold = True
    p.add_run('等待更新').italic = True
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['任务','负责人','状态']):
        cell.text = text
    for values in [('记忆治理','测试角色甲','进行中'),('审批接线','测试角色乙','完成')]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    doc.add_page_break()
    doc.add_heading('第二页：风险与后续', 1)
    doc.add_paragraph('视觉检查应保留页眉、页脚、表格及两页分页。')
    source = root/'source.docx'
    doc.save(source)
    originals = {source.name:digest(source)}
    edited = Document(source)
    target = next(p for p in edited.paragraphs if '等待更新' in p.text)
    target.runs[1].text = '已经更新，格式仍保留'
    edited.save(root/'edited.docx')
    reread = Document(root/'edited.docx')
    run = next(p for p in reread.paragraphs if '格式仍保留' in p.text)
    assert run.runs[0].bold and run.runs[1].italic
    assert len(reread.tables[0].rows)==3
    assert '验收页眉' in reread.sections[0].header.paragraphs[0].text
    checks.append('docx_run_format_header_footer_table_pagebreak_preserved')

    book = Workbook()
    data = book.active
    data.title = '明细'
    for row in [['类别','项目','数量','单价','金额'],['学习','书籍',2,20,'=C2*D2'],
                ['工作','文具',3,5,'=C3*D3'],['学习','课程',1,30,'=C4*D4']]:
        data.append(row)
    data.freeze_panes = 'A2'
    data.auto_filter.ref = 'A1:E4'
    data.print_options.horizontalCentered = True
    for name, width in [('A',16),('B',20),('C',12),('D',12),('E',14)]:
        data.column_dimensions[name].width = width
    for cell in data[1]:
        cell.font = Font(bold=True)
    for row in data.iter_rows(min_row=2,min_col=4,max_col=5):
        for cell in row:
            cell.number_format = '0.00'
    summary = book.create_sheet('汇总')
    summary.append(['项目','数值'])
    summary.append(['学习合计','=SUMIF(明细!A2:A4,"学习",明细!E2:E4)'])
    summary.append(['全部合计','=SUM(明细!E2:E4)'])
    summary.append(['预算状态','=IF(B3<=100,"预算内","超预算")'])
    summary.append(['验收日期',date(2026,10,17)])
    summary['B5'].number_format = 'yyyy-mm-dd'
    summary.column_dimensions['A'].width = 18
    summary.column_dimensions['B'].width = 20
    chart = BarChart()
    chart.title = '合成金额'
    chart.add_data(Reference(data,min_col=5,min_row=1,max_row=4),titles_from_data=True)
    chart.set_categories(Reference(data,min_col=2,min_row=2,max_row=4))
    data.add_chart(chart,'A7')
    source = root/'source.xlsx'
    book.save(source)
    originals[source.name] = digest(source)
    edited = load_workbook(source)
    edited['明细']['C2'] = 3
    edited.save(root/'edited.xlsx')
    edited.close()
    formula_book = load_workbook(root/'edited.xlsx')
    assert formula_book['汇总']['B2'].value.startswith('=SUMIF(')
    assert formula_book['明细'].freeze_panes=='A2' and len(formula_book['明细']._charts)==1
    formula_book.close()
    recalc = convert(root/'edited.xlsx',root/'recalculated',executable=soffice,format='xlsx')
    values = load_workbook(recalc['path'],data_only=True)
    assert [values['汇总'][f'B{i}'].value for i in (2,3,4)]==[90,105,'超预算']
    assert values['汇总']['B5'].value.date()==date(2026,10,17)
    values.close()
    checks.append('xlsx_crosssheet_sumif_sum_if_date_recalculated_90_105_overbudget')

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = '阶段进度：结构保留'
    table = slide.shapes.add_table(3,2,Inches(.5),Inches(1.4),Inches(8),Inches(1.7)).table
    for row, cells in enumerate([['功能','状态'],['记忆','完成'],['设备','待验']]):
        for col, text in enumerate(cells):
            table.cell(row,col).text = text
    slide.notes_slide.notes_text_frame.text = 'NOTES_SENTINEL：合成讲者备注需保留。'
    slide = slides.slides.add_slide(slides.slide_layouts[5])
    slide.shapes.title.text = '第二页：合成图表'
    dataset = CategoryChartData()
    dataset.categories = ['记忆','行为','工具']
    dataset.add_series('通过项',[5,6,3])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.7),Inches(1.5),
                           Inches(8),Inches(4.5),dataset)
    source = root/'source.pptx'
    slides.save(source)
    originals[source.name] = digest(source)
    edited = Presentation(source)
    edited.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = '阶段进度：已经更新'
    edited.save(root/'edited.pptx')
    reread = Presentation(root/'edited.pptx')
    assert len(reread.slides)==2
    assert 'NOTES_SENTINEL' in reread.slides[0].notes_slide.notes_text_frame.text
    assert any(shape.has_table for shape in reread.slides[0].shapes)
    assert any(shape.has_chart for shape in reread.slides[1].shapes)
    checks.append('pptx_edit_preserves_table_chart_notes_second_slide')

    rendered = {}
    for kind, required in [('docx',['复杂结构','第二页','已经更新']),
                           ('xlsx',['学习合计','超预算','105']),('pptx',['已经更新','合成图表'])]:
        # Render the calculated workbook so the PDF is tested against real caches.
        source = Path(recalc['path']) if kind=='xlsx' else root/f'edited.{kind}'
        result = convert(source,root/f'pdf-{kind}',executable=soffice)
        reader = PdfReader(result['path'])
        text = ''.join(''.join(page.extract_text().split()) for page in reader.pages)
        assert len(reader.pages)>=2, (kind,len(reader.pages))
        images = []
        document = pdfium.PdfDocument(result['path'])
        try:
            for index in range(len(document)):
                page = document[index]
                bitmap = page.render(scale=1.3)
                path = root/f'pdf-{kind}'/f'page-{index+1}.png'
                bitmap.to_pil().save(path)
                bitmap.close();page.close()
                images.append(str(path))
        finally:
            document.close()
        missing = [token for token in required if token not in text]
        ocr_text = ''
        if missing:
            # Some LibreOffice Chinese heading glyphs have non-reading-order
            # PDF text. Verify actual pixels rather than dropping the check.
            ocr_text = ''.join(''.join(recognize(path,provider='auto',detect=True)['text'].split())
                               for path in images)
            assert all(token in ocr_text for token in missing), (kind,missing,ocr_text)
        rendered[kind] = {'pdf':result['path'],'pages':len(reader.pages),'images':images,
                          'pdf_text_order_misses':missing,'ocr_fallback_checked':bool(missing)}
    assert all(digest(root/name)==value for name,value in originals.items())
    checks.append('original_sources_unchanged_and_three_formats_rendered')
    return {'passed':True,'checks':checks,'rendered':rendered,
            'external_model_calls':0,'natural_language_quality':'not_tested',
            'visual_review':'pending_manual_review','original_hashes':originals}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--soffice',required=True)
    parser.add_argument('--out',type=Path,required=True)
    args = parser.parse_args()
    args.out.mkdir(parents=True,exist_ok=False)
    report = {'passed':False}
    try:
        report = exercise(args.out.resolve(),args.soffice)
    except BaseException as error:
        report['error'] = str(error)
        raise
    finally:
        (args.out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
        print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':
    main()
