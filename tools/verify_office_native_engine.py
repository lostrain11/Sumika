"""Exercise the shipped official Office engine, never the host soffice install."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import time
import uuid


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--candidate', type=Path, required=True)
    parser.add_argument('--fixtures', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    candidate = args.candidate.resolve(strict=True)
    fixtures = args.fixtures.resolve(strict=True)
    output = args.output.resolve()/uuid.uuid4().hex
    output.mkdir(parents=True)
    node = candidate/'runtime/node/node.exe'
    cli = candidate/'runtime/dsh/node_modules/@deepseek-ai/libreoffice-kit/lib/cli.js'
    python = candidate/'runtime/office/dependencies/python/python.exe'
    originals = {p.name:digest(p) for p in fixtures.glob('edited.*') if p.is_file()}
    report = {'passed':False, 'candidate':str(candidate), 'calls':[], 'external_model_calls':0,
              'boundary':'Actual bundled official engine and library checks on current Windows; no clean-machine or autonomous model-quality claim.'}

    def call(*arguments, success=True):
        start = time.monotonic()
        result = subprocess.run([str(node),str(cli),*map(str,arguments)],
            capture_output=True,text=True,encoding='utf8',timeout=100)
        report['calls'].append({'argv':list(map(str,arguments)), 'returncode':result.returncode,
            'stdout':result.stdout,'stderr':result.stderr,'seconds':round(time.monotonic()-start,3)})
        assert (result.returncode==0)==success, report['calls'][-1]
        return json.loads(result.stdout) if success else None

    try:
        report['capabilities'] = call('capabilities')
        assert report['capabilities']['runtime']['backend']=='native'
        report['conversions'] = {}
        for kind in ('docx','xlsx','pptx'):
            report['conversions'][kind] = call('convert','--input',fixtures/f'edited.{kind}',
                '--output',output/f'{kind}.pdf','--timeout-ms','60000')
            assert report['conversions'][kind]['backend']=='native'
        report['recalculation'] = call('recalculate','--input',fixtures/'edited.xlsx',
            '--output',output/'recalculated.xlsx','--timeout-ms','60000')
        report['images'] = call('render','--input',fixtures/'edited.docx',
            '--output-dir',output/'docx-pages','--pages','1,2','--timeout-ms','60000')
        script = '''
import json,sys
from pathlib import Path
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image
root=Path(sys.argv[1])
pages={kind:len(PdfReader(root/(kind+'.pdf')).pages) for kind in ('docx','xlsx','pptx')}
assert all(n>=2 for n in pages.values()),pages
book=load_workbook(root/'recalculated.xlsx',data_only=True)
sheet=book['汇总']
values=[sheet[cell].value for cell in ('B2','B3','B4')]
assert values==[90,105,'超预算'],values
book.close()
formula=load_workbook(root/'recalculated.xlsx',data_only=False)
assert all(formula['汇总'][cell].data_type=='f' for cell in ('B2','B3','B4'))
formula.close()
images=[]
for path in sorted((root/'docx-pages').rglob('*.png')):
 with Image.open(path) as image:
  assert image.width>100 and image.height>100
  images.append({'path':str(path),'size':image.size})
assert len(images)==2,images
print(json.dumps({'pdf_pages':pages,'cached_values':values,'formulas_preserved':True,'images':images},ensure_ascii=False))
'''
        result = subprocess.run([str(python),'-X','utf8','-B','-c',script,str(output)],
            capture_output=True,text=True,encoding='utf8',timeout=30)
        assert result.returncode==0,result.stdout+result.stderr
        report['file_checks'] = json.loads(result.stdout)
        bad = output/'invalid.docx'
        bad.write_text('not an Office zip',encoding='utf8')
        call('convert','--input',bad,'--output',output/'invalid.pdf',success=False)
        assert not (output/'invalid.pdf').exists()
        assert all(digest(fixtures/name)==value for name,value in originals.items())
        report['source_hashes_unchanged'] = originals
        report['visual_review'] = 'pending'
        report['passed'] = True
    finally:
        (output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
        print('ARTIFACT',output,'PASSED',report['passed'],flush=True)


if __name__=='__main__':
    main()
